"""Build the starter workbook. Headers only. The Desk menu fills the tabs."""
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import tables

BASE = "https://raw.githubusercontent.com/PatelNicknp/nicks-stock-desk/main/data/"
TABS = [
    "scores", "score_details", "market_today", "sectors", "breadth_history",
    "alerts", "score_changes", "portfolio", "portfolio_summary", "watchlist",
    "prices", "news", "scanner_growth", "scanner_quality", "scanner_dividend", "scanner_wheel",
    "scanner_vcp", "scanner_canslim", "scanner_box", "scanner_flag", "scanner_leaders", "scanner_52w",
    "highs_lows_history",
    "earnings_calendar", "earnings_heatmap", "earnings_cards", "earnings_history",
    "put_call", "option_walls", "analyst_ratings", "institutional_13f", "aaii_sentiment",
    "insider_trading", "headline_sentiment", "xly_xlp",
    "financials_summary", "earnings_tracker", "company_financials", "dividend_scanner",
    "macro_screener", "sector_macro", "historical_pe",
    "breakouts", "trends", "pullbacks", "gainers_losers", "sector_performance", "heatmaps",
    "freshness", "updates_log", "help", "settings", "universe",
    "backtest_summary", "backtest_buckets", "backtest_equity", "backtest_yearly", "backtest_notes",
]


def main():
    wb = openpyxl.Workbook()
    setup = wb.active
    setup.title = "Setup"
    setup["A1"] = "Nick's Stock Desk"
    setup["A2"] = "Research only. This workbook does not place trades."
    setup["A3"] = "These tabs only have headers. After you open this file in Google Sheets, paste apps_script/Refresh.gs and use Desk > Refresh data now. The script copies each file into the tab with the same name. Do not use IMPORTDATA."
    setup["A4"] = "Base URL"
    setup["B4"] = BASE
    setup["A6"] = "Last refresh"
    setup["B6"] = ""
    setup["A7"] = "Status"
    setup["B7"] = "Not loaded yet. Use Desk > Refresh data now."
    setup.column_dimensions["A"].width = 88
    setup.column_dimensions["B"].width = 88
    for name in TABS:
        ws = wb.create_sheet(name)
        for i, col in enumerate(tables.COLS[name], 1):
            ws.cell(1, i, col)
    out = os.path.join(ROOT, "data", "Nick_Stock_Desk_Starter.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
